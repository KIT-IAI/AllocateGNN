"""Step A: close and hash-freeze the active inputs of the margin-criterion revision.

The copy list of ``backup_revision_inputs.py`` (``manifests/source_files.sha256.csv``)
is the reference. ``legacy/`` has left active processing (decision D8) and is
not checked. The script

1. adds the legacy connection-code references (upstream history e05d851 and
   its neighbours) as ``git archive`` extracts under ``code/legacy_reference``;
2. re-hashes every active file under ``inputs/`` and ``code/`` and compares it
   with the copy list, the DataOverview inventories and the boundary provenance;
3. checks that the pricing workbook carries T9/T25 and that the DNO GeoJSON has
   the 14 GSP groups, and that the release code archive is reproducible;
4. writes one timestamped receipt with an explicit PASS/FAIL conclusion plus the
   frozen file list ``manifests/active_inputs_<stamp>.sha256.csv``.

Nothing already in the backup is overwritten.
"""
from __future__ import annotations

import argparse
import csv
from datetime import datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import subprocess
import tempfile
import zipfile

LEGACY_COMMITS = {
    # commit: why it is kept
    "c06b690": "001_cost_surface closing state (static surface bitwise EXACT)",
    "38b6e78": "002_criterion nine-distribution panel closing state",
    "e05d851": "011_scale_crossover and 012_decision_adequacy as first committed",
    "b6f6f68": "012 unit-multiplier fix (AU MVA readings), last 011/012 state",
}
LEGACY_PATHS = ["CaseStudy/3_Experiment/4_CostCriterion", "CaseStudy/3_Experiment/0_common",
                "SpatialGranularity/Placement"]
INACTIVE_PREFIXES = ("legacy/",)
ACTIVE_ROOTS = ("inputs", "code")


def long_path(path: Path) -> str:
    value = str(path.absolute())
    return "\\\\?\\" + value if os.name == "nt" and not value.startswith("\\\\?\\") else value


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(long_path(path), "rb") as stream:
        for block in iter(lambda: stream.read(1 << 22), b""):
            h.update(block)
    return h.hexdigest()


def archive_legacy(upstream: Path, backup: Path) -> list[dict]:
    rows = []
    for short, why in LEGACY_COMMITS.items():
        full = subprocess.check_output(["git", "-C", str(upstream), "rev-parse", short], text=True).strip()
        present = [p for p in LEGACY_PATHS
                   if subprocess.run(["git", "-C", str(upstream), "cat-file", "-e", f"{full}:{p}"]).returncode == 0]
        out = backup / "code" / "legacy_reference" / f"upstream_{full[:12]}_connection_code.zip"
        if not out.exists():
            subprocess.run(["git", "-C", str(upstream), "archive", "--format=zip", f"--output={out}", full, "--", *present],
                           check=True)
        rows.append({"source": f"git:{full}:{'|'.join(present)}", "backup_relative_path": out.relative_to(backup).as_posix(),
                     "category": "legacy_connection_code", "bytes": out.stat().st_size, "sha256": sha256(out),
                     "note": why})
    return rows


def boundary_rows(backup: Path) -> list[dict]:
    folder = backup / "inputs" / "boundaries"
    record = json.loads((folder / "provenance.json").read_text(encoding="utf-8"))
    rows = [{"source": meta.get("service", meta.get("url")), "backup_relative_path": f"inputs/boundaries/{name}",
             "category": "figure_boundaries", "bytes": meta["bytes"], "sha256": meta["sha256"],
             "note": f"{meta['version']}; {meta['licence']}; retrieved {meta['retrieved_utc']}"}
            for name, meta in record.items()]
    rows.append({"source": "scripts/fetch_boundaries.py", "backup_relative_path": "inputs/boundaries/provenance.json",
                 "category": "figure_boundaries", "bytes": (folder / "provenance.json").stat().st_size,
                 "sha256": sha256(folder / "provenance.json"), "note": "download provenance"})
    return rows


def pricing_check(backup: Path) -> dict:
    import openpyxl
    folder = backup / "inputs" / "connection_pricing"
    book = openpyxl.load_workbook(folder / "tnuos_tariffs_2026_27.xlsx", read_only=True)
    zones = json.loads((folder / "dno_zones_20240503.geojson").read_text(encoding="utf-8"))
    names = sorted(f["properties"]["Name"] for f in zones["features"])
    expected = sorted(f"_{c}" for c in "ABCDEFGHJKLMNP")
    result = {"workbook_sheets": book.sheetnames, "has_T9": "T9" in book.sheetnames, "has_T25": "T25" in book.sheetnames,
              "dno_features": len(names), "dno_names": names, "dno_crs": zones.get("crs", {}).get("properties", {}).get("name"),
              "dno_matches_14_gsp_groups": names == expected}
    result["pass"] = bool(result["has_T9"] and result["has_T25"] and result["dno_matches_14_gsp_groups"])
    return result


def release_archive_check(upstream: Path, backup: Path, receipt: dict) -> dict:
    commit = receipt["release_commit"]
    stored = backup / "code" / f"upstream_release_{commit[:12]}.zip"
    with tempfile.TemporaryDirectory() as tmp:
        fresh = Path(tmp) / "fresh.zip"
        subprocess.run(["git", "-C", str(upstream), "archive", "--format=zip", f"--output={fresh}", commit], check=True)
        reproduced = sha256(fresh) == sha256(stored)
    with zipfile.ZipFile(stored) as archive:
        names = set(archive.namelist())
    needed = ["sglib/experiment/connection.py", "sglib/experiment/control_fields.py",
              "sglib/experiment/conditional_bounds.py", "sglib/experiment/connection_observations.py",
              "sglib/experiment/planning_tasks.py", "sglib/core/algorithms/pmedian.py"]
    return {"commit": commit, "archive": stored.relative_to(backup).as_posix(), "reproduced_from_git": reproduced,
            "contains_connection_modules": all(n in names for n in needed), "pass": reproduced and all(n in names for n in needed)}


def main() -> None:
    here = Path(__file__).resolve()
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--backup", type=Path, default=here.parents[2] / "results" / "_backup")
    ap.add_argument("--upstream", type=Path, default=here.parents[5])
    args = ap.parse_args()
    backup, upstream = args.backup.resolve(), args.upstream.resolve()
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    copy_receipt = json.loads((backup / "manifests/backup_receipt.json").read_text(encoding="utf-8"))
    copy_manifest = backup / "manifests/source_files.sha256.csv"
    if sha256(copy_manifest) != copy_receipt["manifest_sha256"]:
        raise SystemExit("copy manifest differs from its receipt")
    reference = {r["backup_relative_path"]: r for r in csv.DictReader(copy_manifest.open(encoding="utf-8"))}
    active_reference = {k: v for k, v in reference.items() if not k.startswith(INACTIVE_PREFIXES)}

    supplements = archive_legacy(upstream, backup) + boundary_rows(backup)
    expected = {k: {"sha256": v["sha256"], "bytes": int(v["bytes"]), "category": v["category"], "source": v["source"]}
                for k, v in active_reference.items()}
    for row in supplements:
        if row["backup_relative_path"] in expected:
            raise SystemExit(f"supplement collides with copy list: {row['backup_relative_path']}")
        expected[row["backup_relative_path"]] = {"sha256": row["sha256"], "bytes": int(row["bytes"]),
                                                 "category": row["category"], "source": row["source"]}

    # os.walk on the \\?\ form: Path.rglob silently skips entries beyond MAX_PATH on Windows.
    on_disk, prefix = {}, long_path(backup)
    for root in ACTIVE_ROOTS:
        for folder, _dirs, files in os.walk(long_path(backup / root)):
            for name in files:
                rel = Path(os.path.join(folder, name)[len(prefix) + 1:]).as_posix()
                on_disk[rel] = backup / rel

    missing = sorted(set(expected) - set(on_disk))
    unlisted = sorted(set(on_disk) - set(expected))
    mismatched, frozen = [], []
    for rel in sorted(set(expected) & set(on_disk)):
        path = on_disk[rel]
        digest, size = sha256(path), os.stat(long_path(path)).st_size
        ok = digest == expected[rel]["sha256"] and size == expected[rel]["bytes"]
        if not ok:
            mismatched.append({"path": rel, "expected": expected[rel]["sha256"], "observed": digest})
        frozen.append({"backup_relative_path": rel, "category": expected[rel]["category"], "bytes": size,
                       "sha256": digest, "source": expected[rel]["source"]})

    inventory_failures = []
    for country in ("1_UK", "2_AU"):
        inventory = json.loads((backup / "inputs/base/results/1_DataOverview" / country / "data_inventory.json").read_text(encoding="utf-8"))
        for item in inventory["artifacts"]:
            rel = f"inputs/base/{item['path']}"
            if rel not in on_disk or sha256(on_disk[rel]) != item["sha256"]:
                inventory_failures.append(rel)
    n_inventory = sum(len(json.loads((backup / "inputs/base/results/1_DataOverview" / c / "data_inventory.json")
                                     .read_text(encoding="utf-8"))["artifacts"]) for c in ("1_UK", "2_AU"))

    pricing = pricing_check(backup)
    release = release_archive_check(upstream, backup, copy_receipt)
    legacy_code = sorted(p.name for p in (backup / "code/legacy_reference").iterdir())

    list_path = backup / "manifests" / f"active_inputs_{stamp}.sha256.csv"
    with list_path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(frozen[0]), lineterminator="\n")
        writer.writeheader(); writer.writerows(frozen)
    supplement_path = backup / "manifests" / f"supplement_sources_{stamp}.csv"
    with supplement_path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(supplements[0]), lineterminator="\n")
        writer.writeheader(); writer.writerows(supplements)

    failures = []
    if missing: failures.append(f"{len(missing)} listed files missing")
    if unlisted: failures.append(f"{len(unlisted)} files on disk without a manifest entry")
    if mismatched: failures.append(f"{len(mismatched)} hash/size mismatches")
    if inventory_failures: failures.append(f"{len(inventory_failures)} DataOverview inventory artifacts fail")
    if not pricing["pass"]: failures.append("pricing sources incomplete")
    if not release["pass"]: failures.append("release code archive not reproducible or incomplete")
    receipt = {
        "schema": "margin_criterion_active_input_freeze_v1", "created_utc": stamp,
        "conclusion": "PASS" if not failures else "FAIL",
        "statement": ("活动输入冻结通过" if not failures else "活动输入冻结未通过") ,
        "failures": failures, "scope": "inputs/ and code/ under results/_backup; legacy/ excluded by D8",
        "reference_copy_manifest": {"path": "manifests/source_files.sha256.csv", "sha256": copy_receipt["manifest_sha256"],
                                    "rows_total": len(reference), "rows_active": len(active_reference),
                                    "rows_legacy_excluded": len(reference) - len(active_reference),
                                    "previous_receipt_status": copy_receipt["status"],
                                    "previous_receipt_failure_scope": "legacy/review_prof presentation (UNC path); out of scope by D8"},
        "supplements": {"path": supplement_path.relative_to(backup).as_posix(), "sha256": sha256(supplement_path), "rows": len(supplements)},
        "files_frozen": len(frozen), "bytes_frozen": sum(r["bytes"] for r in frozen),
        "frozen_list": {"path": list_path.relative_to(backup).as_posix(), "sha256": sha256(list_path)},
        "missing": missing, "unlisted": unlisted, "mismatched": mismatched,
        "dataoverview_inventory": {"artifacts": n_inventory, "failures": inventory_failures},
        "pricing": pricing, "release_code": release, "legacy_code_reference": legacy_code,
        "release_id": copy_receipt["release_id"], "release_commit": copy_receipt["release_commit"],
    }
    out = backup / "manifests" / f"active_input_freeze_{stamp}.json"
    out.write_text(json.dumps(receipt, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps({k: receipt[k] for k in ("conclusion", "failures", "files_frozen", "bytes_frozen")}, ensure_ascii=False))
    print(out)
    if failures:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
